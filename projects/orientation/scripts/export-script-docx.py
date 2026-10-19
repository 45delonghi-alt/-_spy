"""timeline.json から読み上げ用の台本（Word）を書き出す。
使い方: python3 scripts/export-script-docx.py ../../docs/script.docx
"""
import json, sys
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

root = Path(__file__).resolve().parent.parent
tl = json.loads((root / "src/data/timeline.json").read_text(encoding="utf-8"))
out = Path(sys.argv[1] if len(sys.argv) > 1 else root / "../../docs/script.docx")


def ts(sec):
    return f"{int(sec // 60)}:{int(sec % 60):02d}"


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Mm(297), Mm(210)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Mm(15))
st = doc.styles["Normal"]
st.font.name = "Yu Gothic"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")

doc.add_heading(tl["meta"]["title"].replace("【", "\n【") + "　台本", level=1)
doc.add_paragraph(
    "動画尺 6:00（停止時間は含まず）。「停止」欄のある場面は、ナレーション終了後に一時停止して討議します。"
)

widths = [Mm(20), Mm(38), Mm(55), Mm(154)]
for ch in tl["chapters"]:
    doc.add_heading(
        f"CH{ch['id']:02d}　{ch['missionTitle']}（{ch['informationTheme']}）　{ts(ch['start'])}–{ts(ch['end'])}",
        level=2,
    )
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, h in enumerate(["時間", "場面", "画面メッセージ", "ナレーション"]):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "142A43")
    for s in [x for x in tl["scenes"] if x["chapter"] == ch["id"]]:
        r = t.add_row().cells
        r[0].text = f"{ts(s['start'])}–\n{ts(s['end'])}"
        r[1].text = f"SC{s['id']:02d}\n{s['title']}"
        r[2].text = s.get("keyMessage", "")
        r[3].text = s.get("narration", "")
        q = s.get("pause")
        if q:
            p = r[3].add_paragraph()
            run = p.add_run(f"【停止 {q['id']}】{q['question']}（{q['guide']}）")
            run.bold = True
            run.font.color.rgb = RGBColor(0xDC, 0x09, 0x24)
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = w
        for c in t.columns[i].cells:
            c.width = w

out.parent.mkdir(parents=True, exist_ok=True)
doc.save(out)
print(out)
